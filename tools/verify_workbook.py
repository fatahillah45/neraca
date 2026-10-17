"""Verifikasi workbook Neraca secara independen dengan openpyxl.

Pemakaian: python tools/verify_workbook.py <folder berisi workbook dan expected.json>
Memeriksa: nama sheet, nama dan rentang Excel Table, kolom, jumlah baris, format Rupiah,
serta angka yang harus sama dengan hasil mesin hitung (expected.json).
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

RP = '"Rp"#,##0;-"Rp"#,##0'
failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def table_rows(ws, table):
    """Baris data (tanpa header dan total) sebagai list of dict."""
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    header = [ws.cell(min_row, c).value for c in range(min_col, max_col + 1)]
    last = max_row - (1 if table.totalsRowCount else 0)
    rows = []
    for r in range(min_row + 1, last + 1):
        rows.append({h: ws.cell(r, c).value for h, c in zip(header, range(min_col, max_col + 1))})
    return header, rows, (min_col, min_row, max_col, max_row)


def main(folder):
    folder = Path(folder)
    exp = json.loads((folder / 'expected.json').read_text(encoding='utf-8'))

    wb = load_workbook(folder / exp['weeklyFile'])
    check(wb.sheetnames == exp['sheets'], f"sheet: {wb.sheetnames} != {exp['sheets']}")

    summary = wb['Summary']
    check('Daily_Summary' in summary.tables, 'Daily_Summary table missing')
    check('Account_Balances' in summary.tables, 'Account_Balances table missing')
    _, srows, _ = table_rows(summary, summary.tables['Daily_Summary'])
    check(len(srows) == 7, f'Daily_Summary has {len(srows)} rows, expected 7')
    for col, key in (('Income', 'income'), ('Expense', 'expense'), ('Human being', 'giving'), ('Invested (buy)', 'investBuy'), ('Surplus', 'surplus'), ('Transactions', 'count')):
        total = sum((r.get(col) or 0) for r in srows)
        check(total == exp['week'][key], f'Summary {col}: {total} != {exp["week"][key]}')
    _, brows, _ = table_rows(summary, summary.tables['Account_Balances'])
    end = {r['Account']: r['Balance Sun 23:59'] for r in brows}
    check(end == exp['balancesEnd'], f'saldo akhir: {end} != {exp["balancesEnd"]}')

    names = set()
    for day in exp['days']:
        ws = wb[day['sheet']]
        check(len(ws.tables) == 1, f"{day['sheet']}: harus tepat 1 tabel, ada {len(ws.tables)}")
        tname = next(iter(ws.tables))
        table = ws.tables[tname]
        check(tname not in names, f'nama tabel ganda {tname}')
        names.add(tname)
        header, rows, (min_col, min_row, _, _) = table_rows(ws, table)
        check(header == exp['dayColumns'], f"{day['sheet']}: kolom {header}")
        real = [r for r in rows if r.get('ID')]
        check(len(real) == day['count'], f"{day['sheet']}: {len(real)} rows, expected {day['count']}")
        net = sum(r['Cash effect (Rp)'] for r in real)
        check(net == day['netCash'], f"{day['sheet']}: efek kas {net} != {day['netCash']}")
        check(table.totalsRowCount == 1, f"{day['sheet']}: baris total tidak ada")
        check(ws.freeze_panes == 'A5', f"{day['sheet']}: freeze panes {ws.freeze_panes}")
        if real:
            fmt = ws.cell(min_row + 1, min_col + 9).number_format
            check(fmt == RP, f"{day['sheet']}: format Jumlah {fmt!r}")
        else:
            check(ws['A2'].value == 'No transactions.', f"{day['sheet']}: empty-day note missing")

    wa = load_workbook(folder / exp['allFile'])
    check(wa.sheetnames == ['Transactions', 'Accounts', 'Portfolio', 'Info'], f'all-export sheets: {wa.sheetnames}')
    ws = wa['Transactions']
    header, rows, _ = table_rows(ws, ws.tables['All_Transactions'])
    check(header == exp['allColumns'], f'kolom semua: {header}')
    check(len(rows) == exp['allCount'], f'semua transaksi: {len(rows)} != {exp["allCount"]}')
    check(len({r['ID'] for r in rows}) == len(rows), 'ID transaksi ganda di ekspor semua')

    if failures:
        print('FAILED:')
        for f in failures:
            print(' -', f)
        sys.exit(1)
    print(f"PASSED: {exp['weeklyFile']} ({len(wb.sheetnames)} sheet) dan {exp['allFile']} ({len(rows)} transaksi)")


if __name__ == '__main__':
    main(sys.argv[1])
